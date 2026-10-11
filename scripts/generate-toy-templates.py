#!/usr/bin/env python3
"""Generate original, editable TOY Office presentation starter templates."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

OUT = Path("src-tauri/templates")
STYLES = {
    "kurumsal": ("Kurumsal Sunum", (27, 54, 83), (36, 145, 177)),
    "kosgeb-proje": ("Proje / KOSGEB Sunumu", (28, 78, 67), (83, 169, 131)),
    "egitim": ("Eğitim Sunumu", (72, 52, 105), (153, 117, 207)),
}

def rgb(value):
    return RGBColor(*value)

def text_box(slide, left, top, width, height, text, size, color, bold=False):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Liberation Sans"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)
    return shape

def generate(name, title, primary, accent):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for page in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb((250, 251, 252))
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(.22))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(accent)
        bar.line.fill.background()
        if page == 0:
            text_box(slide, .9, 1.9, 11.6, 1.1, title, 36, primary, True)
            text_box(slide, .9, 3.15, 11.5, .7, "Başlık ve sunum bilgilerini düzenleyin", 19, (86, 96, 107))
        elif page == 1:
            text_box(slide, .9, .85, 11.5, .9, "Gündem", 32, primary, True)
            for i, label in enumerate(("Amaç ve kapsam", "Yaklaşım", "Sonraki adımlar")):
                text_box(slide, 1.05, 2.05 + i * 1.05, 10.8, .6, f"{i + 1:02d}  {label}", 22, primary)
        else:
            text_box(slide, .9, .85, 11.5, .9, "İçerik", 32, primary, True)
            text_box(slide, 1.05, 2.15, 11.1, 2.8, "Bu alanı kendi metinleriniz, tablolarınız ve görsellerinizle değiştirin.", 23, (86, 96, 107))
        text_box(slide, .9, 6.95, 11, .3, f"TOY Templates  •  {page + 1}", 10, (112, 122, 132))
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / f"toy-{name}.pptx"
    prs.save(target)
    return target

if __name__ == "__main__":
    from zipfile import ZipFile
    for name, (title, primary, accent) in STYLES.items():
        path = generate(name, title, primary, accent)
        with ZipFile(path) as archive:
            assert archive.testzip() is None
            assert "[Content_Types].xml" in archive.namelist()
        print(f"Generated {path}")
